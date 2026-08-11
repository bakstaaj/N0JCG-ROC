#!/usr/bin/env python3
"""Build the branded N0JCG Gateway DOCX from the canonical Markdown guide."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
VERSION = (ROOT / "VERSION").read_text(encoding="utf-8").strip()
SOURCE = ROOT / "docs" / f"N0JCG_Gateway_End_User_Guide_v{VERSION}.md"
OUTPUT = ROOT / "docs" / "publications" / f"N0JCG_Gateway_End_User_Guide_v{VERSION}.docx"
LOGO = ROOT / "web" / "assets" / "N0JCG_Header_Dark_Approved.png"
ICON = ROOT / "web" / "assets" / "N0JCG_Icon_Approved.png"

NAVY = "0A1F44"
BLUE = "1565C0"
CYAN = "00B8D9"
SLATE = "2B3440"
MIST = "F4F7FA"
MUTED = "536171"
WHITE = "FFFFFF"
BORDER = "C7CDD4"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_font(run, *, name="Arial", size=10.5, color=SLATE, bold=False, italic=False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def paragraph_border(paragraph, color=BORDER, size="6") -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    properties.append(borders)


def page_field(paragraph) -> None:
    for field_type, value in (("begin", None), (None, " PAGE "), ("separate", None)):
        run = paragraph.add_run()._r
        if field_type:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), field_type)
        else:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = value
        run.append(node)
    placeholder = paragraph.add_run("1")
    set_font(placeholder, size=8, color=MUTED)
    end_run = paragraph.add_run()._r
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)


def picture_alt(shape, title: str, description: str) -> None:
    shape._inline.docPr.set("title", title)
    shape._inline.docPr.set("descr", description)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(SLATE)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, SLATE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    if ICON.exists():
        icon = header.add_run().add_picture(str(ICON), width=Inches(0.22))
        picture_alt(icon, "N0JCG platform icon", "Approved N0JCG Open Radio Platform icon.")
    lead = header.add_run(f"  N0JCG  /  GATEWAY END USER GUIDE  /  v{VERSION}")
    set_font(lead, size=8.2, color=BLUE, bold=True)
    paragraph_border(header)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    lead = footer.add_run("N0JCG Open Radio Platform  |  Page ")
    set_font(lead, size=8, color=MUTED)
    page_field(footer)


def add_cover(doc: Document) -> None:
    for _ in range(2):
        doc.add_paragraph()
    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner.autofit = False
    cell = banner.cell(0, 0)
    cell.width = Inches(6.45)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    shade(cell, NAVY)
    cell_margins(cell, top=300, start=300, bottom=300, end=300)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = p.add_run().add_picture(str(LOGO), width=Inches(4.95))
    picture_alt(logo, "N0JCG primary logo", "Approved N0JCG Open Radio Platform logo.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    r = p.add_run("OPERATOR HANDBOOK")
    set_font(r, size=9, color=CYAN, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("N0JCG Gateway")
    set_font(r, size=29, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Installation, configuration, operation, and recovery")
    set_font(r, size=12, color=MUTED)

    metadata = doc.add_table(rows=2, cols=4)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.style = "Table Grid"
    values = (("RELEASE", VERSION, "PUBLICATION", "August 2026"),
              ("PLATFORM", "Ubuntu 24.04 LTS", "AUDIENCE", "Operators and maintainers"))
    for row_index, values_row in enumerate(values):
        for column, value in enumerate(values_row):
            cell = metadata.cell(row_index, column)
            shade(cell, MIST)
            cell_margins(cell, top=120, start=100, bottom=120, end=100)
            run = cell.paragraphs[0].add_run(value)
            set_font(run, size=8.5, color=NAVY, bold=column % 2 == 0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("Open radio systems, engineered as one platform.")
    set_font(r, size=11.5, color=SLATE, italic=True)
    doc.add_page_break()


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            set_font(paragraph.add_run(text[cursor:match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), bold=True)
        elif token.startswith("`"):
            set_font(paragraph.add_run(token[1:-1]), name="Consolas", size=9.2, color=NAVY)
        else:
            set_font(paragraph.add_run(token[1:-1]), italic=True)
        cursor = match.end()
    if cursor < len(text):
        set_font(paragraph.add_run(text[cursor:]))


def add_table(doc: Document, lines: list[str]) -> None:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            properties = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            properties.append(repeat)
        for column in range(columns):
            cell = row.cells[column]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell)
            shade(cell, NAVY if row_index == 0 else (MIST if row_index % 2 == 0 else WHITE))
            value = values[column] if column < len(values) else ""
            run = cell.paragraphs[0].add_run(value.replace("**", "").replace("`", ""))
            set_font(run, size=8.7, color=WHITE if row_index == 0 else SLATE, bold=row_index == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_contents(doc: Document, lines: list[str]) -> None:
    doc.add_paragraph("Contents", style="Heading 1")
    for line in lines:
        match = re.match(r"^## (\d+)\. (.+)$", line)
        if not match:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        number = p.add_run(match.group(1).zfill(2) + "    ")
        set_font(number, size=9.3, color=CYAN, bold=True)
        title = p.add_run(match.group(2))
        set_font(title, size=9.3, color=NAVY, bold=True)
        paragraph_border(p, BORDER, "4")
    doc.add_page_break()


def add_body(doc: Document, lines: list[str]) -> None:
    start = next(index for index, line in enumerate(lines) if line.startswith("## 1. "))
    index = start
    while index < len(lines):
        line = lines[index].rstrip()
        if not line or line == "---":
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            shade(cell, "EEF2F6")
            cell_margins(cell, top=100, start=140, bottom=100, end=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run("\n".join(code_lines))
            set_font(run, name="Consolas", size=8.2, color=NAVY)
            if language:
                run.font.name = "Consolas"
            index += 1
            continue
        if line.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:])
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:])
        elif line.startswith("> "):
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            shade(cell, MIST)
            cell_margins(cell, top=120, start=180, bottom=120, end=180)
            p = cell.paragraphs[0]
            add_inline(p, line[2:])
            p.runs[0].font.color.rgb = rgb(NAVY)
        elif line.startswith("- "):
            parts = [line[2:]]
            index += 1
            while index < len(lines):
                candidate = lines[index].rstrip()
                if (not candidate or candidate.startswith(("#", "|", "> ", "- ", "```"))
                        or re.match(r"^\d+\. ", candidate)):
                    break
                parts.append(candidate.strip())
                index += 1
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, " ".join(parts))
            continue
        elif re.match(r"^\d+\. ", line):
            match = re.match(r"^(\d+)\. (.*)", line)
            assert match is not None
            item_number = match.group(1)
            parts = [match.group(2)]
            index += 1
            while index < len(lines):
                candidate = lines[index].rstrip()
                if (not candidate or candidate.startswith(("#", "|", "> ", "- ", "```"))
                        or re.match(r"^\d+\. ", candidate)):
                    break
                parts.append(candidate.strip())
                index += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.375)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            marker = p.add_run(f"{item_number}.  ")
            set_font(marker, size=10.5, color=BLUE, bold=True)
            add_inline(p, " ".join(parts))
            continue
        else:
            paragraph_lines = [line]
            index += 1
            while index < len(lines):
                candidate = lines[index].rstrip()
                if (not candidate or candidate == "---" or candidate.startswith(("#", "|", "> ", "- ", "```"))
                        or re.match(r"^\d+\. ", candidate)):
                    break
                paragraph_lines.append(candidate)
                index += 1
            p = doc.add_paragraph()
            add_inline(p, " ".join(part.strip() for part in paragraph_lines))
            continue
        index += 1


def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.core_properties.title = f"N0JCG Gateway End User Guide v{VERSION}"
    doc.core_properties.subject = "Installation, configuration, operation, and recovery for N0JCG Gateway"
    doc.core_properties.author = "N0JCG Open Radio Platform"
    doc.core_properties.keywords = "N0JCG, Gateway, ROC, APRS, Winlink, weather, Amateur Radio"
    doc.core_properties.comments = "Generated from the canonical Markdown source."
    configure(doc)
    add_cover(doc)
    add_contents(doc, lines)
    add_body(doc, lines)
    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
